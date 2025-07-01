# utils.py

import fitz  # PyMuPDF
import docx
from fpdf import FPDF
import io

def extract_text_from_pdf(uploaded_file):
    text = ""
    with fitz.open(stream=uploaded_file.read(), filetype="pdf") as doc:
        for page in doc:
            text += page.get_text()
    return text.strip()

def extract_text_from_docx(uploaded_file):
    doc = docx.Document(uploaded_file)
    return "\n".join([para.text for para in doc.paragraphs]).strip()

class StyledPDF(FPDF):
    def header(self):
        self.set_font("Arial", "B", 14)
        self.cell(0, 10, "Resume", ln=True, align="C")
        self.ln(5)

    def add_section(self, title, content):
        self.set_font("Arial", "B", 12)
        self.set_text_color(0, 0, 128)
        self.cell(0, 10, title, ln=True)
        self.set_text_color(0, 0, 0)
        self.set_font("Arial", "", 11)
        self.multi_cell(0, 8, content)
        self.ln(3)

def generate_pdf_resume(content, filename="resume.pdf"):
    """
    Generates a stylized multi-page PDF resume.
    Returns the PDF as a BytesIO buffer for download.
    """
    pdf = StyledPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)

    # Split content by sections marked with ###
    sections = content.strip().split("###")
    for section in sections:
        if ":" in section:
            title, body = section.strip().split(":", 1)
            pdf.add_section(title.strip(), body.strip())
        else:
            pdf.add_section("Details", section.strip())

    buffer = io.BytesIO()
    pdf.output(buffer)
    buffer.seek(0)
    return buffer
